"""Real HTTP/database/export coverage using synthetic deals, never live data."""
import io
import zipfile

import fitz
import openpyxl
import pytest

pytestmark = pytest.mark.asyncio


async def new_deal(client, name):
    response = await client.post("/api/deals", json={"project_name": name, "property_type": "multifamily"})
    assert response.status_code == 200
    return response.json()["id"]


@pytest.mark.parametrize("endpoint", ["compare", "compare/export"])
async def test_comparison_rejects_duplicate_missing_and_deleted_deals(client, endpoint):
    a, b = await new_deal(client, "A"), await new_deal(client, "B")
    url = "/api/deals/" + endpoint
    assert (await client.post(url, json={"deal_ids": [a, a]})).status_code == 400
    assert (await client.post(url, json={"deal_ids": [a, 99999]})).status_code == 404
    await client.delete(f"/api/deals/{b}")
    assert (await client.post(url, json={"deal_ids": [a, b]})).status_code == 404


async def test_return_values_agree_in_detail_compare_and_exports(client):
    a, b = await new_deal(client, "A & B <Fund>"), await new_deal(client, "Sale Deal")
    metrics = {
        "target_returns": {"primary_strategy": "hold", "target_irr": 22,
            "hold_scenario": {"cash_on_cash_return": 8},
            "sale_scenario": {"sale_irr": 22, "sale_equity_multiple": 2.4, "is_hypothetical": True}},
        "deal_structure": {"ltv": 60},
    }
    await client.put(f"/api/deals/{a}", json={"metrics": metrics, "scores": {"overall": None}, "notes": "A < B & C"})
    await client.put(f"/api/deals/{b}", json={"metrics": {"target_returns": {"target_irr": 15}, "deal_structure": {"ltv": 80}}})
    detail = (await client.get(f"/api/deals/{a}")).json()
    assert detail["target_irr"] is None
    assert detail["target_equity_multiple"] is None
    comparison = (await client.post("/api/deals/compare", json={"deal_ids": [b, a]})).json()["deals"]
    assert [d["id"] for d in comparison] == [b, a]
    assert comparison[1]["target_irr"] is None
    xlsx = await client.post("/api/deals/compare/export", json={"deal_ids": [b, a]})
    assert xlsx.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
    rows = {row[0].value: row for row in sheet.iter_rows() if row[0].value}
    assert rows["Target IRR"][2].value is None
    assert rows["LTV"][1].fill.fill_type is None  # High leverage is not a winner.
    pdf = await client.get(f"/api/reports/deal/{a}/pdf")
    assert pdf.status_code == 200
    with fitz.open(stream=pdf.content, filetype="pdf") as document:
        text = "".join(page.get_text() for page in document)
    assert "22.00%" not in text
    assert "8.00%" in text
    assert "Not scored" in text
    assert "A < B & C" in text


async def test_document_csv_upload_read_reprocess_and_delete(client):
    deal = await new_deal(client, "Document journey")
    uploaded = await client.post(f"/api/deals/{deal}/documents/upload",
        files={"file": ("facts.csv", b"Metric,Value\nUnits,120\nEquity,2000000\n", "text/csv")})
    assert uploaded.status_code == 200
    doc_id = uploaded.json()["id"]
    docs = await client.get(f"/api/deals/{deal}/documents")
    assert docs.status_code == 200
    assert any(d["id"] == doc_id for d in docs.json())
    file = await client.get(f"/api/deals/documents/{doc_id}/file")
    assert file.status_code == 200 and b"Units,120" in file.content
    reprocessed = await client.post(f"/api/deals/documents/{doc_id}/reprocess")
    assert reprocessed.status_code == 200
    assert reprocessed.json()["text_length_after"] > 0
    assert (await client.delete(f"/api/deals/documents/{doc_id}")).status_code == 200
    assert (await client.get(f"/api/deals/documents/{doc_id}/file")).status_code == 404


async def test_portfolio_exports_and_chat_history(client):
    deal = await new_deal(client, "Export coverage")
    investment = await client.post("/api/investments/", json={
        "project_name": "Export coverage", "amount_invested": 10000,
        "investment_date": "2025-01-01", "deal_id": deal,
    })
    assert investment.status_code == 200
    inv_id = investment.json()["id"]
    distribution = await client.post(f"/api/investments/{inv_id}/distributions",
        json={"date": "2025-12-31", "amount": 500})
    assert distribution.status_code == 200
    for path in ["/api/investments/portfolio", "/api/investments/portfolio/analytics",
                 f"/api/investments/{inv_id}/performance", "/api/deals/pipeline/summary",
                 f"/api/chat/history/{deal}"]:
        response = await client.get(path)
        assert response.status_code == 200, (path, response.text)
    assert (await client.delete(f"/api/chat/history/{deal}")).status_code == 200
    workbook = await client.get("/api/reports/portfolio/excel")
    assert workbook.status_code == 200
    assert openpyxl.load_workbook(io.BytesIO(workbook.content)).sheetnames
    pdf = await client.get("/api/reports/portfolio/quarterly/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    exported = await client.get("/api/reports/export/json")
    assert exported.status_code == 200
    assert exported.json()
    csv = await client.get("/api/reports/export/csv")
    assert csv.status_code == 200
    assert zipfile.ZipFile(io.BytesIO(csv.content)).namelist()
